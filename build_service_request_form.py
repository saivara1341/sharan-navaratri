from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path('/Users/saivaraprasad/Downloads/Siddhi_Dynamics_LLP_Service_Order_Form_Personal_Account.docx')
OUTPUT = Path('/Users/saivaraprasad/Downloads/siddhidynamics-main/Siddhi_Dynamics_LLP_Service_Order_and_Client_Request_Form.docx')


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=55, start=100, bottom=55, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, color='D9D9D9', sz='8'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), sz)
        element.set(qn('w:color'), color)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(int(inches * 1440)))
    tc_w.set(qn('w:type'), 'dxa')


def format_run(run, bold=False, size=9, color='222222', italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')


def write_cell(cell, text, bold=False, size=9, color='222222', shade=None):
    cell.text = ''
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)
    if shade:
        set_cell_shading(cell, shade)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    format_run(r, bold=bold, size=size, color=color)


def add_section_bar(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_width(cell, 7.05)
    write_cell(cell, text, bold=True, size=10.5, color='FFFFFF', shade='1C1816')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return table


def add_two_column_fields(doc, fields):
    table = doc.add_table(rows=len(fields), cols=2)
    table.autofit = False
    for row, (label, blank) in zip(table.rows, fields):
        set_width(row.cells[0], 2.15)
        set_width(row.cells[1], 4.9)
        write_cell(row.cells[0], label, bold=True, shade='F4F1ED')
        write_cell(row.cells[1], blank)
    return table


def add_spacer(doc, points=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(points)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    return p


def add_body(doc, text, size=9, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    format_run(r, size=size, italic=italic)
    return p


doc = Document(SOURCE)

doc.add_section(WD_SECTION.NEW_PAGE)
section = doc.sections[-1]
section.top_margin = Inches(0.35)
section.bottom_margin = Inches(0.35)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(4)
r = p.add_run('CLIENT SERVICE REQUEST AND ONBOARDING FORM')
format_run(r, bold=True, size=14, color='1C1816')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(10)
r = p.add_run('Complete and submit this page before work starts. It helps us confirm the scope, collect the required materials, and schedule your service correctly.')
format_run(r, size=8.5, color='555555')

add_section_bar(doc, '1   SERVICE REQUEST DETAILS')
add_two_column_fields(doc, [
    ('Client name / business', '________________________________________________________________________________'),
    ('Primary contact', 'Name: ____________________   Phone / email: _______________________________________'),
    ('Requested service', '________________________________________________________________________________'),
    ('Requested start date', '______________________   Preferred contact time: _________________________________'),
])
add_spacer(doc)

add_section_bar(doc, '2   PROJECT REQUIREMENTS')
add_two_column_fields(doc, [
    ('Business goal', 'What result do you need from this service? ____________________________________________'),
    ('Target audience', '________________________________________________________________________________'),
    ('Key requirements', '________________________________________________________________________________\n________________________________________________________________________________'),
    ('Reference links / examples', '________________________________________________________________________________'),
])
add_spacer(doc)

add_section_bar(doc, '3   MATERIALS AND ACCESS TO PROVIDE')
add_body(doc, 'Tick what you are providing now. If something is not ready, write the expected date. Please do not write passwords on this form; provide access securely by invitation or through an agreed secure method.', size=8, italic=True)
materials = doc.add_table(rows=4, cols=2)
materials.autofit = False
items = [
    ('☐ Logo / brand files', '☐ Website / domain / hosting access'),
    ('☐ Content, images, or product details', '☐ Social-media / ad-account access'),
    ('☐ Existing website or system details', '☐ Analytics / CRM / WhatsApp access'),
    ('☐ Other: __________________________', 'Expected date for pending items: __________________'),
]
for row, (left, right) in zip(materials.rows, items):
    set_width(row.cells[0], 3.52)
    set_width(row.cells[1], 3.53)
    write_cell(row.cells[0], left, size=8.5)
    write_cell(row.cells[1], right, size=8.5)
add_spacer(doc)

add_section_bar(doc, '4   SUBMISSION AND CONFIRMATION')
add_body(doc, 'How to submit: Complete both pages, sign the Service Order and Payment Agreement, then send the filled form to Siddhi Dynamics LLP by WhatsApp or email at the contact shown on page 1. We will confirm receipt and the service start date after reviewing your request, required materials, access, and agreed payment.', size=8.5)
confirm = doc.add_table(rows=2, cols=1)
confirm.autofit = False
set_width(confirm.cell(0, 0), 7.05)
set_width(confirm.cell(1, 0), 7.05)
write_cell(confirm.cell(0, 0), '☐ I confirm that the information above is correct and that I will provide any remaining materials or access by the stated date.', size=8.5)
write_cell(confirm.cell(1, 0), 'Client name: ______________________________   Signature: ______________________________   Date: _______________', size=8.5)

add_body(doc, 'Service start condition: Siddhi Dynamics LLP will begin work after the service request is complete, the scope is confirmed, the required materials or access are received (or a delivery date is agreed), and the agreed advance payment has been received.', size=8, italic=True)

doc.save(OUTPUT)
print(OUTPUT)
